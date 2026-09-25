#!/usr/bin/env python3
"""One-off FM/PM structural audit (analysis only). Outputs xlsx + markdown summary."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "biological-targets"
OUT_XLSX = ROOT.parent / "FM-PM-Structural-Audit.xlsx"
OUT_MD = ROOT.parent / "FM-PM-Structural-Audit-Summary.md"

BRS_DIRS = [f"brs{i}" for i in range(1, 7)]


def load_frontmatter(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---"):
        return {}
    end = text.index("---", 3)
    block = text[3:end]
    return yaml.safe_load(block) or {}


def condense(text: str | None, max_len: int = 420) -> str:
    if not text:
        return ""
    t = re.sub(r"\s+", " ", str(text)).strip()
    if len(t) <= max_len:
        return t
    return t[: max_len - 1].rstrip() + "…"


def phenome_list(data: dict) -> str:
    rels = data.get("phenome_relationships") or []
    parts = []
    for r in rels:
        label = r.get("target_phenome") or r.get("phenome_label") or r.get("phenome_id")
        if not label:
            continue
        rt = r.get("relationship_type")
        parts.append(f"{label} ({rt})" if rt else str(label))
    return "; ".join(parts) if parts else "None mapped in PM front matter"


def scan() -> tuple[list[dict], list[dict]]:
    fms: list[dict] = []
    pms: list[dict] = []
    for brs_dir in BRS_DIRS:
        brs_path = DOCS / brs_dir
        if not brs_path.is_dir():
            continue
        brs = brs_dir.upper()
        for fm_folder in sorted(brs_path.iterdir()):
            if not fm_folder.is_dir() or not re.match(r"fm\d", fm_folder.name):
                continue
            for mdx in sorted(fm_folder.glob("*.mdx")):
                data = load_frontmatter(mdx)
                is_pm = bool(data.get("pm_id")) or re.search(r"-pm\d", mdx.name, re.I)
                if is_pm:
                    pms.append(
                        {
                            "brs": brs,
                            "path": mdx,
                            "pm_id": data.get("pm_id"),
                            "title": data.get("title", ""),
                            "parent_fm": data.get("parent_fm", ""),
                            "mission": data.get("mission"),
                            "summary": data.get("summary"),
                            "phenomes": phenome_list(data),
                        }
                    )
                else:
                    fms.append(
                        {
                            "brs": brs,
                            "fm_id": data.get("fm_id", ""),
                            "title": data.get("title", ""),
                            "mission": data.get("mission"),
                            "summary": data.get("summary"),
                            "mechanisms_covered": data.get("mechanisms_covered") or [],
                        }
                    )
    fms.sort(key=lambda x: x["fm_id"])
    pms.sort(key=lambda x: x["pm_id"] or "")
    return fms, pms


# --- Analytical judgments (repository-informed; not new ontology) ---

FM_JUDGMENTS: dict[str, dict] = {
    "BRS1(FM1)": {
        "config": "MIXED (B + A)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Four-step supply chain: substrate → LAT1 transport → distinct monoamine signalling arms. PM3/PM4 are testable signalling nodes but interpret best with upstream PM1–PM2 context.",
    },
    "BRS1(FM2)": {
        "config": "E",
        "pm_isolated": "Yes",
        "fm_operationalised": "Partly",
        "integrated_state": "Partly",
        "integrated_pm": "BRS1-FM2-PM5 — Acetylcholine Synthesis Support",
        "notes": "Single PM covers essentially the whole FM proposition (ACh synthesis). FM title is broader cholinergic function than this one conversion step.",
    },
    "BRS1(FM3)": {
        "config": "E",
        "pm_isolated": "Yes",
        "fm_operationalised": "Partly",
        "integrated_state": "Partly",
        "integrated_pm": "BRS1-FM3-PM6 — Neuronal Membrane DHA Incorporation",
        "notes": "FM names membrane composition/fluidity; sole PM is DHA incorporation — structural lipid scope wider than one PM.",
    },
    "BRS1(FM4)": {
        "config": "MIXED (A + D)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "BRS1-FM4-PM7 — GABA–Glutamate Neurotransmission Balance",
        "notes": "Decomposed synthesis/clearance/excitotoxicity nodes plus PM7 as emergent E/I balance state overlapping FM mission.",
    },
    "BRS2(FM1)": {
        "config": "B",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "BRS2-FM1-PM4 — Methionine Cycle Flux",
        "notes": "Classic cyclic pathway: parallel remethylation routes → SAMe → integrated flux readout. PM4 explicitly models whole-cycle coordination.",
    },
    "BRS2(FM2)": {
        "config": "B",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Sequential transsulfuration diversion then glutathione synthesis; FM is coupling concept spanning both PMs.",
    },
    "BRS2(FM3)": {
        "config": "E",
        "pm_isolated": "Yes",
        "fm_operationalised": "Partly",
        "integrated_state": "Partly",
        "integrated_pm": "BRS2-FM3-PM7 — Phosphatidylcholine Formation",
        "notes": "FM is cross-system coupling (methylation→membrane); single PM is one biochemical bridge.",
    },
    "BRS3(FM1)": {
        "config": "C",
        "pm_isolated": "Yes",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Parallel upstream (NF-κB) and interface (gut-derived) inflammatory entry points converging on tone.",
    },
    "BRS3(FM2)": {
        "config": "MIXED (A + C)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "BRS3-FM2-PM4 — ROS Generation vs Clearance Balance",
        "notes": "Decomposed antioxidant modules (Nrf2, recycling, lipid protection) plus net ROS balance integrator.",
    },
    "BRS3(FM3)": {
        "config": "MIXED (C + D)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Unclear",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "FM mission emphasises pro-resolution (SPM); PM7 is downstream cytokine modulation — integrated resolution state not fully decomposed. Flag: FM proposition vs PM7 emphasis mismatch.",
    },
    "BRS4(FM1)": {
        "config": "A",
        "pm_isolated": "Yes",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Three separable bioenergetic components (ETC, NAD+, creatine buffer) under one capacity umbrella.",
    },
    "BRS4(FM2)": {
        "config": "C",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "ROS generation/control vs structural mitochondrial protection — parallel contributors to resilience.",
    },
    "BRS4(FM3)": {
        "config": "MIXED (B + A + D)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "BRS4-FM3-PM8 — Metabolic Fuel Switching",
        "notes": "Transport/utilisation steps plus PM8 as integrator of multi-substrate flexibility.",
    },
    "BRS4(FM4)": {
        "config": "E",
        "pm_isolated": "Yes",
        "fm_operationalised": "Partly",
        "integrated_state": "Partly",
        "integrated_pm": "BRS4-FM4-PM9 — Mitochondrial Biogenesis",
        "notes": "Biogenesis PM is core of expansion/adaptation FM; broader adaptive phenotype not further decomposed.",
    },
    "BRS5(FM1)": {
        "config": "MIXED (A + C)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Barrier structure, endotoxin containment, and keystone ecology — parallel with partial sequential dependency (barrier before LPS load).",
    },
    "BRS5(FM2)": {
        "config": "MIXED (B + C)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Ecological turnover enables metabolite signalling arms (SCFA, polyphenol metabolites).",
    },
    "BRS5(FM3)": {
        "config": "C",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Vagal/ENS signalling vs gut-side precursor biotransformation — parallel gut–brain routes.",
    },
    "BRS6(FM1)": {
        "config": "B",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Clear meal-level sequence: glucose appearance → variability → insulin-mediated disposal.",
    },
    "BRS6(FM2)": {
        "config": "MIXED (B + C)",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Entrainment scaffold (PM5) and cortisol rhythm output (PM4) — mutually dependent timing pair.",
    },
    "BRS6(FM3)": {
        "config": "C",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Sympathetic–parasympathetic shift vs vagal/HRV readout — related but separable autonomic lenses.",
    },
    "BRS6(FM4)": {
        "config": "C",
        "pm_isolated": "Partly",
        "fm_operationalised": "Partly",
        "integrated_state": "Yes",
        "integrated_pm": "No",
        "notes": "Chronic metabolic-inflammatory signalling vs stress-linked appetite/reward — parallel load-allocation axes.",
    },
}


PM_JUDGMENTS: dict[str, dict] = {
    "BRS1-FM4-PM7": {
        "effects": "Net excitatory–inhibitory balance between glutamate and GABA signalling.",
        "isolatable": "Partly",
        "measurable": "Partly",
        "sibling_dep": "Yes",
        "fm_integrated": "Yes",
        "notes": "Repository text defines PM7 as net balance, not isolated synthesis/clearance — overlaps FM-level E/I state.",
    },
    "BRS2-FM1-PM4": {
        "effects": "Integrated methionine-cycle throughput and routing (methionine→SAM→homocysteine→remethylation/transsulfuration).",
        "isolatable": "Partly",
        "measurable": "Partly",
        "sibling_dep": "Yes",
        "fm_integrated": "Yes",
        "notes": "Flux PM aggregates sibling steps; homocysteine is junction readout not sole objective.",
    },
    "BRS4-FM3-PM8": {
        "effects": "Cross-substrate metabolic flexibility (glucose/fat/ketone switching).",
        "isolatable": "Partly",
        "measurable": "Partly",
        "sibling_dep": "Yes",
        "fm_integrated": "Yes",
        "notes": "Explicit integrator across PM6/PM7 transport and utilisation pathways.",
    },
    "BRS3-FM2-PM4": {
        "effects": "Net ROS production vs clearance balance.",
        "isolatable": "Partly",
        "measurable": "Partly",
        "sibling_dep": "Yes",
        "fm_integrated": "Yes",
        "notes": "System-level redox balance rather than single enzyme/pathway.",
    },
    "BRS3-FM3-PM7": {
        "effects": "Downstream cytokine network intensity (IL-6, TNF-α, CRP-linked tone).",
        "isolatable": "Partly",
        "measurable": "Partly",
        "sibling_dep": "Partly",
        "fm_integrated": "No",
        "notes": "Under resolution FM but PM targets cytokine modulation — not identical to SPM-centric FM mission.",
    },
}


def default_pm_judgment(pm: dict, fm_j: dict) -> dict:
    pid = pm["pm_id"]
    if pid in PM_JUDGMENTS:
        return PM_JUDGMENTS[pid]
    n = len([p for p in scan_pms_cache if p.get("parent_fm") == pm["parent_fm"]])
    cfg = fm_j.get("config", "")
    sibling = "Partly" if n > 1 and "B" in cfg or "MIXED" in cfg else "No"
    if n == 1:
        sibling = "No"
    return {
        "effects": condense(pm.get("mission") or pm.get("summary"), 180),
        "isolatable": "Yes" if n == 1 or "A" in cfg else "Partly",
        "measurable": "Partly",
        "sibling_dep": sibling,
        "fm_integrated": "No",
        "notes": "Default: atomic mechanism proposition as authored; see FM configuration for cluster context.",
    }


REL_JUDGMENTS: dict[str, dict] = {
    "BRS1-FM1-PM1": {
        "character": "sequential stage",
        "contribute": "Substrate ceiling for all downstream monoamine pathways.",
        "gap": "Transport competition, receptor-level signalling, and pathway-specific conversion.",
        "direct_fm": "Partly",
        "why": "Protein adequacy trials rarely isolate FM-level monoaminergic function without transport/signalling confounds.",
        "caution": "Do not treat amino-acid intake studies as full FM efficacy without PM2–PM4 context.",
    },
    "BRS1-FM1-PM2": {
        "character": "sequential stage",
        "contribute": "Blood–brain barrier precursor delivery bias after meals.",
        "gap": "Absolute precursor pool and post-synaptic signalling outcomes.",
        "direct_fm": "Partly",
        "why": "LAT1 manipulation is mechanistically direct for PM2; FM outcome needs signalling PMs.",
        "caution": "Carbohydrate:protein ratio evidence maps to PM2, not whole monoaminergic FM.",
    },
    "BRS1-FM1-PM3": {
        "character": "component",
        "contribute": "Noradrenergic arousal/executive modulation arm.",
        "gap": "Serotonin arm, precursor supply, LAT1 context.",
        "direct_fm": "Partly",
        "why": "Catecholamine-focused evidence is direct for PM3; FM integrates dual monoamine axes.",
        "caution": "Avoid inferring complete monoaminergic FM from NA-only endpoints.",
    },
    "BRS1-FM1-PM4": {
        "character": "component",
        "contribute": "Serotonergic regulation arm.",
        "gap": "Precursor transport and NA arm.",
        "direct_fm": "Partly",
        "why": "5-HT pathway evidence is PM-direct; FM requires coordinated multi-PM interpretation.",
        "caution": "Tryptophan manipulation ≠ full monoaminergic FM without PM1–PM2.",
    },
    "BRS1-FM4-PM7": {
        "character": "integrator",
        "contribute": "Emergent E/I balance state across glutamate and GABA arms.",
        "gap": "Does not alone specify synthesis, clearance, or excitotoxic buffering mechanisms.",
        "direct_fm": "Partly",
        "why": "Human GABA/glutamate balance markers are closer to FM mission than sibling atomic PMs.",
        "caution": "PM7 evidence may be best FM-facing; still distinguish from dietary lever claims on PM8.",
    },
    "BRS1-FM4-PM8": {
        "character": "component",
        "contribute": "GAD-dependent GABA synthesis capacity from glutamate.",
        "gap": "Net E/I balance, clearance, excitotoxicity integration.",
        "direct_fm": "No",
        "why": "GAD/PLP biology is PM-specific; human brain synthesis flux not established.",
        "caution": "Strong PM mechanistic evidence ≠ FM-level GABA balance improvement.",
    },
    "BRS2-FM1-PM4": {
        "character": "integrator",
        "contribute": "Whole-cycle flux coordination and routing under load.",
        "gap": "Individual remethylation enzyme steps when studied in isolation.",
        "direct_fm": "Partly",
        "why": "Homocysteine/SAM markers partially reflect FM efficiency but conflate parallel PM1–PM3.",
        "caution": "Cycle-flux readouts aggregate multiple PMs — avoid double-counting at FM rollup.",
    },
    "BRS3-FM3-PM7": {
        "character": "contributor",
        "contribute": "Downstream cytokine tone modulation.",
        "gap": "SPM/eicosanoid resolution biology emphasised in FM mission.",
        "direct_fm": "No",
        "why": "Cytokine endpoints do not fully test pro-resolution FM proposition.",
        "caution": "FM resolution narrative must not be inferred solely from PM7 cytokine evidence.",
    },
    "BRS3-FM3-PM8": {
        "character": "component",
        "contribute": "Pro-resolving vs pro-inflammatory lipid mediator balance.",
        "gap": "Cytokine-network modulation and upstream transcription.",
        "direct_fm": "Partly",
        "why": "Omega-3/SPM evidence is relatively direct for this PM and closest to FM resolution theme.",
        "caution": "SPM PM evidence still not equivalent to integrated resolution capacity without PM7 context.",
    },
    "BRS4-FM3-PM8": {
        "character": "integrator",
        "contribute": "Integrated metabolic fuel-switching capacity.",
        "gap": "Single-pathway transport or ketone oxidation in isolation.",
        "direct_fm": "Partly",
        "why": "Flexibility metrics (e.g. respiratory quotient shifts) align with FM3 more than one transport PM.",
        "caution": "Ketogenic/flexibility interventions may touch PM7/PM8 — attribute evidence by mechanism tested.",
    },
    "BRS6-FM1-PM1": {
        "character": "sequential stage",
        "contribute": "Upstream glucose entry profile shaping post-prandial load.",
        "gap": "Insulin disposal and cognitive energy outcomes without variability context.",
        "direct_fm": "Partly",
        "why": "GI/glycaemic index evidence is direct for appearance kinetics; FM needs PM2–PM3.",
        "caution": "Meal-structure evidence ≠ full glycaemic–cognitive FM without downstream PMs.",
    },
    "BRS6-FM1-PM2": {
        "character": "sequential stage",
        "contribute": "Volatility/oscillation control after glucose entry.",
        "gap": "Disposal efficiency and cognitive endpoints alone.",
        "direct_fm": "Partly",
        "why": "CGM variability metrics map well to PM2; FM integrates stability + disposal.",
        "caution": "Do not roll variability trials into insulin sensitivity claims without PM3.",
    },
    "BRS6-FM1-PM3": {
        "character": "sequential stage",
        "contribute": "Insulin-mediated clearance and tissue uptake.",
        "gap": "Meal composition effects on appearance kinetics.",
        "direct_fm": "Partly",
        "why": "Clamp/HOMA-style evidence is PM3-direct; FM cognitive-energy claim needs full cascade.",
        "caution": "Insulin sensitivity studies are indirect for glucose appearance PM1.",
    },
}


def default_rel_judgment(pm: dict, fm_j: dict) -> dict:
    pid = pm["pm_id"]
    if pid in REL_JUDGMENTS:
        return REL_JUDGMENTS[pid]
    cfg = fm_j.get("config", "C")
    if cfg == "E":
        char = "near-equivalent"
    elif "B" in cfg and "sequential" not in cfg:
        char = "sequential stage"
    elif "A" in cfg:
        char = "component"
    else:
        char = "contributor"
    return {
        "character": char,
        "contribute": condense(pm.get("mission"), 160),
        "gap": "Sibling PM mechanisms and FM-level integrated state not captured by this PM alone.",
        "direct_fm": "Partly" if char != "near-equivalent" else "Partly",
        "why": "PM-authored mission is narrower than FM summary/mission; typical studies target PM-scale biology.",
        "caution": "Surface PM findings at FM only with explicit aggregation rules; distinguish visibility from FM-level proof.",
    }


scan_pms_cache: list[dict] = []


def fm_label(fm: dict) -> str:
    return f"{fm['fm_id']} — {fm['title']}"


def pm_label(pm: dict) -> str:
    return f"{pm['pm_id']} — {pm['title']}"


def constituent_pms(fm: dict) -> str:
    parts = []
    for m in fm.get("mechanisms_covered") or []:
        parts.append(f"{m.get('id')} — {m.get('name')}")
    return "; ".join(parts)


def write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(14, len(headers[col - 1]) * 0.9))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(fms: list[dict], pms: list[dict]) -> openpyxl.Workbook:
    global scan_pms_cache
    scan_pms_cache = pms
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    fm_rows = []
    for fm in fms:
        j = FM_JUDGMENTS.get(fm["fm_id"], {})
        desc = condense(fm.get("mission") or fm.get("summary"))
        fm_rows.append(
            [
                fm["brs"],
                fm_label(fm),
                desc,
                constituent_pms(fm),
                len(fm.get("mechanisms_covered") or []),
                j.get("config", "Unclear"),
                j.get("pm_isolated", "Unclear"),
                j.get("fm_operationalised", "Unclear"),
                j.get("integrated_state", "Unclear"),
                j.get("integrated_pm", "No"),
                j.get("notes", ""),
            ]
        )

    pm_rows = []
    rel_rows = []
    fm_by_id = {f["fm_id"]: f for f in fms}
    for pm in pms:
        fm = fm_by_id.get(pm["parent_fm"], {})
        fj = FM_JUDGMENTS.get(pm["parent_fm"], {})
        pj = default_pm_judgment(pm, fj)
        desc = condense(pm.get("mission") or pm.get("summary"))
        pm_rows.append(
            [
                pm["brs"],
                fm_label(fm) if fm else pm["parent_fm"],
                pm_label(pm),
                desc,
                pj.get("effects", ""),
                pj.get("isolatable", "Unclear"),
                pj.get("measurable", "Unclear"),
                pj.get("sibling_dep", "Unclear"),
                pj.get("fm_integrated", "Unclear"),
                pm.get("phenomes", ""),
                pj.get("notes", ""),
            ]
        )
        rj = default_rel_judgment(pm, fj)
        rel_rows.append(
            [
                pm["brs"],
                fm_label(fm) if fm else pm["parent_fm"],
                condense(fm.get("mission") or fm.get("summary")) if fm else "",
                pm_label(pm),
                desc,
                rj.get("character", ""),
                rj.get("contribute", ""),
                rj.get("gap", ""),
                rj.get("direct_fm", "Unclear"),
                rj.get("why", ""),
                rj.get("caution", ""),
            ]
        )

    write_sheet(
        wb,
        "FM Overview",
        [
            "BRS",
            "FM",
            "FM Description",
            "Constituent PMs",
            "Number of PMs",
            "Proposed FM/PM Configuration",
            "Can constituent PMs be experimentally isolated?",
            "Can the FM itself be operationalised/measured independently of individual PMs?",
            "Does the FM represent an integrated biological state/capacity?",
            "Does any PM appear to represent an integrated FM-level state rather than an atomic mechanism?",
            "Notes / rationale",
        ],
        fm_rows,
    )
    write_sheet(
        wb,
        "PM Detail",
        [
            "BRS",
            "Parent FM",
            "PM",
            "PM Description",
            "Primary Biological Effect(s)",
            "PM experimentally isolatable?",
            "PM has a distinct measurable biological proposition?",
            "PM depends on sibling PMs for its biological interpretation?",
            "PM appears to represent an integrated FM state?",
            "Existing Phenome relationships",
            "Notes / rationale",
        ],
        pm_rows,
    )
    write_sheet(
        wb,
        "FM x PM Relationship",
        [
            "BRS",
            "FM",
            "FM Description",
            "PM",
            "PM Description",
            "Relationship Character",
            "What does this PM contribute to the FM?",
            "What remains biologically unrepresented if this PM is considered alone?",
            "Is PM evidence likely to be direct evidence for the FM proposition?",
            "Why?",
            "Evidence Aggregation Caution",
        ],
        rel_rows,
    )
    return wb


def write_summary(fms: list[dict], pms: list[dict]) -> None:
    by_brs_fm: dict[str, int] = {}
    by_brs_pm: dict[str, int] = {}
    for fm in fms:
        by_brs_fm[fm["brs"]] = by_brs_fm.get(fm["brs"], 0) + 1
    for pm in pms:
        by_brs_pm[pm["brs"]] = by_brs_pm.get(pm["brs"], 0) + 1

    config_counts: dict[str, int] = {}
    for fm in fms:
        c = FM_JUDGMENTS.get(fm["fm_id"], {}).get("config", "Unclear")
        config_counts[c] = config_counts.get(c, 0) + 1

    single_pm = [fm_label(f) for f in fms if len(f.get("mechanisms_covered") or []) == 1]

    lines = [
        "# FM–PM Structural Audit Summary (BRS1–BRS6)",
        "",
        "Analysis-only pass over current repository front matter (21 FMs, 51 PMs). No ontology or page edits.",
        "",
        "## Counts by BRS",
        "",
        "| BRS | FMs | PMs |",
        "|-----|-----|-----|",
    ]
    for brs in [f"BRS{i}" for i in range(1, 7)]:
        lines.append(f"| {brs} | {by_brs_fm.get(brs, 0)} | {by_brs_pm.get(brs, 0)} |")
    lines.append(f"| **Total** | **{len(fms)}** | **{len(pms)}** |")
    lines += [
        "",
        "## Configuration distribution (Sheet 1)",
        "",
    ]
    for k, v in sorted(config_counts.items(), key=lambda x: -x[1]):
        lines.append(f"- **{k}:** {v} FM(s)")
    lines += [
        "",
        "## Clearest exemplars",
        "",
        "- **A (Decomposed):** BRS4(FM1) Cellular Bioenergetics — ETC, NAD⁺, creatine buffer as separable modules.",
        "- **B (Pathway/network):** BRS2(FM1) Methylation Cycle Efficiency; BRS6(FM1) glycaemic cascade (appearance → variability → disposal).",
        "- **C (Parallel contributors):** BRS3(FM1) anti-inflammatory tone; BRS6(FM4) metabolic inflammation vs stress-appetite axes.",
        "- **D (Integrative/emergent):** BRS1(FM4) with PM7 as E/I balance; BRS4(FM3) PM8 fuel switching integrator.",
        "- **E (Near-equivalent / single-PM FMs):** BRS1(FM2), BRS1(FM3), BRS2(FM3), BRS4(FM4).",
        "- **MIXED:** BRS1(FM1), BRS1(FM4), BRS3(FM2), BRS3(FM3), BRS4(FM3), BRS5(FM1), BRS5(FM2), BRS6(FM2).",
        "",
        "## Additional configuration pattern (flagged)",
        "",
        "**Supply-chain / layered dependency:** monoaminergic and glycaemic FMs where PM order matters for interpretation (not merely parallel). Treated under MIXED B or explicit sequential relationship rows.",
        "",
        "## Structural anomalies (flagged, not fixed)",
        "",
        f"- **Single-PM FMs ({len(single_pm)}):** " + "; ".join(single_pm) + ".",
        "- **PM duplicates or encodes FM-level state:** BRS1-FM4-PM7 (E/I balance); BRS2-FM1-PM4 (cycle flux); BRS4-FM3-PM8 (fuel switching); BRS3-FM2-PM4 (net ROS balance).",
        "- **FM mission vs PM emphasis mismatch:** BRS3(FM3) resolution-capacity FM vs BRS3-FM3-PM7 cytokine modulation focus.",
        "- **FM broader than sole PM:** BRS1(FM2–FM3), BRS2(FM3), BRS4(FM4) titles exceed single-PM scope.",
        "- **Sibling-dependent interpretation clusters:** BRS1(FM1), BRS1(FM4), BRS2(FM1), BRS4(FM3), BRS6(FM1).",
        "",
        "## Implications for future PM → FM evidence interpretation",
        "",
        "- **Membership ≠ FM proof:** PM phenome mappings and PM findings may be surfaced at FM level without constituting FM-level evidence.",
        "- **Integrator PMs** need different rollup rules than atomic component PMs.",
        "- **Sequential FMs** require ordered or conditional aggregation (downstream PM evidence assumes upstream context).",
        "- **Single-PM FMs** collapse structural distinction but still separate authorship layers (FM integrated narrative vs PM mechanism evidence).",
        "- **Parallel FMs** need disjunctive/convergent synthesis, not simple vote counting across PMs.",
        "",
        "## Which distinct FM/PM structural patterns must future evidence-aggregation methodology accommodate?",
        "",
        "1. **Decomposed multi-PM modules (A)** — independent PM evidence with FM synthesis across non-redundant components.",
        "2. **Pathway / sequential networks (B)** — stage-dependent evidence chains where upstream PM results constrain downstream interpretation.",
        "3. **Parallel contributors (C)** — partial, non-substitutable evidence lanes merging into one FM capacity.",
        "4. **Emergent / integrator PMs within FMs (D)** — PMs that approximate FM-level states (E/I balance, cycle flux, fuel switching, ROS balance).",
        "5. **Near-equivalent single-PM FMs (E)** — FM rollup largely tracks one PM but must preserve FM-vs-PM evidence typing.",
        "6. **Mixed configurations** — explicit multi-label FMs (common in BRS1, BRS3, BRS4, BRS5, BRS6).",
        "7. **Supply-chain dependency** — monoaminergic and glycaemic stacks requiring contextual, not additive, PM→FM lifts.",
        "8. **FM–PM proposition mismatch** — aggregation must not treat PM endpoints as proxies for a materially wider (or different) FM mission.",
        "",
    ]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    fms, pms = scan()
    wb = build_workbook(fms, pms)
    wb.save(OUT_XLSX)
    write_summary(fms, pms)
    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_MD}")


if __name__ == "__main__":
    main()
